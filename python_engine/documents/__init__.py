"""Document Engine — reusable document capabilities.

Capabilities:
    - Create, edit, analyze documents
    - Insert tables, images, indexes, headers, footers
    - Create automated reports
    - Extract content and structure
    - Template-based generation
"""

from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import datetime
import os

CORPORATE_BLUE_DARK = "1F4E79"
CORPORATE_BLUE = "2E75B6"
CORPORATE_BLUE_LIGHT = "D9E2F3"
CORPORATE_WHITE = "FFFFFF"


class DocumentEngine:
    """Reusable document capabilities. NOT a tool — a library."""

    def create_document(self, path: str = None) -> dict:
        """Create a new empty document."""
        doc = Document()
        if path:
            doc.save(path)
        return {"path": path} if path else {"created": True}

    def open_document(self, path: str) -> dict:
        """Open an existing document and return structure info."""
        doc = Document(path)
        info = {
            "path": path,
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "sections": len(doc.sections),
            "size": os.path.getsize(path),
        }
        return info

    def add_heading(self, path: str, text: str, level: int = 1) -> dict:
        """Add a heading to the document."""
        doc = Document(path)
        doc.add_heading(text, level=level)
        doc.save(path)
        return {"path": path, "heading": text, "level": level}

    def add_paragraph(self, path: str, text: str,
                      bold: bool = False, italic: bool = False,
                      size: int = 11, alignment: str = "left",
                      color: str = None) -> dict:
        """Add a paragraph with formatting."""
        align_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        doc = Document(path)
        para = doc.add_paragraph()
        para.alignment = align_map.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*self._parse_color(color))
        doc.save(path)
        return {"path": path, "paragraphs": len(doc.paragraphs)}

    def insert_table(self, path: str, headers: list = None,
                     rows: list = None, style: str = "Table Grid") -> dict:
        """Insert a table with headers and data."""
        doc = Document(path)
        ncols = len(headers) if headers else (len(rows[0]) if rows else 1)
        nrows = (1 if headers else 0) + len(rows or [])
        table = doc.add_table(rows=nrows, cols=ncols)
        table.style = style
        if headers:
            for i, h in enumerate(headers):
                table.rows[0].cells[i].text = str(h)
        for ri, row in enumerate(rows or [], start=1 if headers else 0):
            for ci, val in enumerate(row):
                table.rows[ri].cells[ci].text = str(val)
        doc.save(path)
        stats = {"path": path, "rows": nrows, "cols": ncols}
        return stats

    def insert_image(self, path: str, image_path: str,
                     width: float = None, add_paragraph: bool = True) -> dict:
        """Insert an image into the document."""
        doc = Document(path)
        if add_paragraph and len(doc.paragraphs) > 0:
            doc.add_paragraph()
        if width:
            doc.add_picture(image_path, width=Inches(width))
        else:
            doc.add_picture(image_path)
        doc.save(path)
        return {"path": path, "image": image_path}

    def add_header(self, path: str, text: str) -> dict:
        """Add text to the document header."""
        doc = Document(path)
        section = doc.sections[0]
        header = section.header
        para = header.paragraphs[0]
        para.text = text
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.save(path)
        return {"path": path, "header": text}

    def add_footer(self, path: str, text: str) -> dict:
        """Add text to the document footer."""
        doc = Document(path)
        section = doc.sections[0]
        footer = section.footer
        para = footer.paragraphs[0]
        para.text = text
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.save(path)
        return {"path": path, "footer": text}

    def add_page_break(self, path: str) -> dict:
        """Add a page break."""
        doc = Document(path)
        doc.add_page_break()
        doc.save(path)
        return {"path": path}

    def create_report(self, path: str, title: str = "Report",
                      sections: list = None, author: str = None,
                      subtitle: str = None, date: str = None,
                      header: str = None, footer: str = None,
                      include_toc: bool = True,
                      page_size: str = "letter") -> dict:
        """Create a corporate report: cover, TOC field, styled sections/tables."""
        doc = Document()
        section = doc.sections[0]
        self._apply_page_setup(section, page_size)
        self._apply_header_footer(section, header or title, footer)

        cover = doc.add_heading(title, 0)
        cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._tint_runs(cover, CORPORATE_BLUE_DARK)
        if subtitle:
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = sub.add_run(subtitle)
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        stamp = date or datetime.date.today().strftime("%B %d, %Y")
        meta_bits = [stamp] + (["Prepared by: " + author] if author else [])
        meta = doc.add_paragraph(" — ".join(meta_bits))
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()

        if include_toc:
            toc_heading = doc.add_heading("Table of Contents", level=1)
            self._tint_runs(toc_heading, CORPORATE_BLUE_DARK)
            self._add_toc_field(doc)

        for section_item in sections or []:
            stype = section_item.get("type", "text")
            heading = section_item.get("heading", "")
            content = section_item.get("content", "")
            if heading:
                level = max(1, min(3, int(section_item.get("level", 1))))
                hd = doc.add_heading(heading, level=level)
                self._tint_runs(hd, CORPORATE_BLUE_DARK if level == 1 else CORPORATE_BLUE)
            if stype == "text":
                for line in content.split("\n"):
                    if line.strip():
                        doc.add_paragraph(line)
            elif stype == "table":
                self._add_styled_table(doc, section_item.get("headers", []),
                                       section_item.get("rows", []))
            quote = section_item.get("quote", "")
            if quote:
                qp = doc.add_paragraph()
                qp.paragraph_format.left_indent = Inches(0.5)
                run = qp.add_run(quote)
                run.italic = True
                run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

        doc.save(path)
        return {"path": path, "title": title, "sections": len(sections or [])}

    @staticmethod
    def _tint_runs(paragraph, hex_color: str) -> None:
        rgb = RGBColor(*DocumentEngine._parse_color(hex_color))
        for run in paragraph.runs:
            run.font.color.rgb = rgb

    @staticmethod
    def _apply_page_setup(section, page_size: str) -> None:
        if str(page_size).lower() == "a4":
            section.page_width, section.page_height = Mm(210), Mm(297)
        else:
            section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.5)

    def _apply_header_footer(self, section, header: str, footer: str) -> None:
        if header:
            para = section.header.paragraphs[0]
            para.text = ""
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(header)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        para = section.footer.paragraphs[0]
        para.text = ""
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer:
            run = para.add_run(footer + "   ")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        run = para.add_run("Page ")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        self._add_page_field(para)

    @staticmethod
    def _add_page_field(paragraph) -> None:
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        num = paragraph.add_run("1")
        num.font.size = Pt(9)
        num.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        run_end = paragraph.add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run_end._r.append(fld_end)

    @staticmethod
    def _add_toc_field(doc) -> None:
        para = doc.add_paragraph()
        run = para.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = 'TOC \\o "1-3" \\h \\z \\u'
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        para.add_run("Table of contents — right-click and choose Update Field.")
        run_end = para.add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run_end._r.append(fld_end)

    def _add_styled_table(self, doc, headers: list, rows: list):
        """Corporate table: blue header with white bold text, banded rows."""
        rows = [list(r) for r in (rows or [])]
        ncols = len(headers) if headers else (len(rows[0]) if rows else 1)
        norm = []
        for row in rows:
            cells = [str(v) for v in row[:ncols]]
            cells += [""] * (ncols - len(cells))
            norm.append(cells)
        table = doc.add_table(rows=(1 if headers else 0) + len(norm), cols=ncols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if headers:
            tr = table.rows[0]._tr
            trPr = tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
            for i, h in enumerate(headers):
                self._set_cell(table.rows[0].cells[i], str(h),
                               bold=True, color=CORPORATE_WHITE,
                               fill=CORPORATE_BLUE_DARK)
        for ri, row in enumerate(norm, start=1 if headers else 0):
            fill = CORPORATE_BLUE_LIGHT if (ri - (1 if headers else 0)) % 2 == 0 else None
            for ci, val in enumerate(row):
                self._set_cell(table.rows[ri].cells[ci], val, fill=fill)
        return table

    @staticmethod
    def _set_cell(cell, value: str, bold: bool = False,
                  color: str = None, fill: str = None) -> None:
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(value)
        run.bold = bold
        run.font.size = Pt(11)
        if color:
            run.font.color.rgb = RGBColor(*DocumentEngine._parse_color(color))
        if fill:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), fill)
            tcPr.append(shd)

    def extract_content(self, path: str) -> dict:
        """Extract all text content and structure from a document."""
        doc = Document(path)
        content = []
        for para in doc.paragraphs:
            content.append({
                "text": para.text,
                "style": para.style.name if para.style else None,
                "bold": any(r.bold for r in para.runs),
            })
        tables = []
        for table in doc.tables:
            data = []
            for row in table.rows:
                data.append([cell.text for cell in row.cells])
            tables.append(data)
        return {"path": path, "paragraphs": content, "tables": tables}

    def analyze_document(self, path: str) -> dict:
        """Analyze document structure and statistics."""
        doc = Document(path)
        word_count = 0
        char_count = 0
        for para in doc.paragraphs:
            word_count += len(para.text.split())
            char_count += len(para.text)
        sections_info = []
        for section in doc.sections:
            sections_info.append({
                "page_width": section.page_width,
                "page_height": section.page_height,
                "orientation": "landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "portrait",
            })
        return {
            "path": path,
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "sections": len(doc.sections),
            "word_count": word_count,
            "char_count": char_count,
            "sections_info": sections_info,
        }

    @staticmethod
    def _parse_color(hex_color: str) -> tuple:
        hex_color = hex_color.lstrip("#")
        if len(hex_color) == 6:
            return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        return (0, 0, 0)

    def mail_merge(self, template_path: str, output_dir: str,
                   data_csv: str = None, rows: list = None,
                   output_pattern: str = "doc_{index}.docx") -> dict:
        """Simple CSV->DOCX mail merge: replaces {{field}} placeholders.

        Args:
            template_path: DOCX with {{field}} placeholders in text/tables.
            output_dir: where to write the merged files.
            data_csv: CSV file with a header row (field names).
            rows: alternative to data_csv, a list of dicts.
            output_pattern: filename pattern with {index} and field names.
        """
        import csv as _csv
        records = list(rows or [])
        if data_csv:
            with open(data_csv, newline="", encoding="utf-8-sig") as f:
                records.extend(list(_csv.DictReader(f)))
        if not records:
            return {"status": "error", "error": "No merge data: provide data_csv or rows"}
        os.makedirs(output_dir, exist_ok=True)
        files = []
        for i, record in enumerate(records, start=1):
            doc = Document(template_path)
            for para in doc.paragraphs:
                self._merge_paragraph(para, record)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            self._merge_paragraph(para, record)
            for section in doc.sections:
                for para in section.header.paragraphs:
                    self._merge_paragraph(para, record)
                for para in section.footer.paragraphs:
                    self._merge_paragraph(para, record)
            name = output_pattern.format(index=i, **{k: str(v) for k, v in record.items()})
            out = os.path.join(output_dir, name)
            doc.save(out)
            files.append(out)
        return {"template": template_path, "count": len(files), "files": files}

    @staticmethod
    def _merge_text(text: str, record: dict) -> str:
        for key, value in record.items():
            text = text.replace("{{" + str(key) + "}}", str(value))
        return text

    @classmethod
    def _merge_paragraph(cls, para, record: dict) -> None:
        # Merge on the whole paragraph text so {{field}} placeholders split
        # across runs still resolve; first-run formatting is preserved.
        if "{{" not in para.text:
            return
        merged = cls._merge_text(para.text, record)
        if merged == para.text:
            return
        if para.runs:
            para.runs[0].text = merged
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.add_run(merged)
