"""Presentation Engine — reusable slide deck capabilities.

Capabilities:
    - Create, edit presentations
    - Add, reorder, delete slides
    - Insert charts, images, tables
    - Apply themes and layouts
    - Create complete automated presentations
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os


class PresentationEngine:
    """Reusable presentation capabilities. NOT a tool — a library."""

    def create_presentation(self, path: str = None) -> dict:
        """Create a new empty presentation."""
        prs = Presentation()
        if path:
            prs.save(path)
        return {"path": path} if path else {"created": True}

    def open_presentation(self, path: str) -> dict:
        """Open an existing presentation and return its structure."""
        prs = Presentation(path)
        info = {
            "path": path,
            "slides": len(prs.slides),
            "slide_width": prs.slide_width,
            "slide_height": prs.slide_height,
            "slide_layouts": [layout.name for layout in prs.slide_layouts],
            "size": os.path.getsize(path),
        }
        return info

    def add_slide(self, path: str, layout: str = None,
                  title: str = None, content: str = None) -> dict:
        """Add a new slide with optional title and content."""
        prs = Presentation(path)
        layout_map = {l.name: l for l in prs.slide_layouts}
        slide_layout = layout_map.get(layout) if layout else prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        if title and slide.shapes.title:
            slide.shapes.title.text = title
        if content:
            body_shape = slide.shapes.placeholders[1] if len(slide.placeholders) > 1 else None
            if body_shape:
                tf = body_shape.text_frame
                tf.clear()
                for line in content.split("\n"):
                    if line.strip():
                        p = tf.add_paragraph()
                        p.text = line
        prs.save(path)
        return {"path": path, "slide": len(prs.slides), "title": title}

    def add_text_slide(self, path: str, title: str,
                       body_lines: list = None) -> dict:
        """Add a title-and-content slide."""
        prs = Presentation(path)
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        if slide.shapes.title:
            slide.shapes.title.text = title
        body = slide.placeholders[1]
        tf = body.text_frame
        for line in body_lines or []:
            p = tf.add_paragraph()
            p.text = line
            p.level = 0
        prs.save(path)
        return {"path": path, "slide": len(prs.slides)}

    def add_table_slide(self, path: str, title: str,
                        headers: list = None, rows: list = None) -> dict:
        """Add a slide with a table."""
        prs = Presentation(path)
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        if title and slide.shapes.title:
            slide.shapes.title.text = title
        nrows = (1 if headers else 0) + len(rows or [])
        ncols = len(headers) if headers else (len(rows[0]) if rows else 1)
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(8)
        height = Inches(0.4 * nrows)
        table = slide.shapes.add_table(nrows, ncols, left, top, width, height).table
        if headers:
            for i, h in enumerate(headers):
                table.cell(0, i).text = str(h)
                for para in table.cell(0, i).text_frame.paragraphs:
                    para.font.bold = True
        for ri, row in enumerate(rows or [], start=1 if headers else 0):
            for ci, val in enumerate(row):
                table.cell(ri, ci).text = str(val)
        prs.save(path)
        return {"path": path, "rows": nrows, "cols": ncols}

    def add_image_slide(self, path: str, image_path: str,
                        title: str = None) -> dict:
        """Add a slide with an image."""
        prs = Presentation(path)
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        if title and slide.shapes.title:
            slide.shapes.title.text = title
        left = Inches(1)
        top = Inches(1.5)
        slide.shapes.add_picture(image_path, left, top)
        prs.save(path)
        return {"path": path, "image": image_path}

    def set_slide_title(self, path: str, slide_index: int,
                        title: str) -> dict:
        """Set the title of a specific slide."""
        prs = Presentation(path)
        if slide_index < len(prs.slides):
            slide = prs.slides[slide_index]
            if slide.shapes.title:
                slide.shapes.title.text = title
        prs.save(path)
        return {"path": path, "slide_index": slide_index}

    def delete_slide(self, path: str, slide_index: int) -> dict:
        """Delete a slide by index."""
        prs = Presentation(path)
        if slide_index < len(prs.slides):
            slide = prs.slides[slide_index]
            rId = prs.slides._sldIdLst[slide_index].get(qn("r:id"))
            prs.part.drop_rel(rId)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[slide_index])
        prs.save(path)
        return {"path": path, "deleted_index": slide_index}

    def add_notes(self, path: str, slide_index: int = 0,
                  notes: str = "") -> dict:
        """Add speaker notes to a slide."""
        prs = Presentation(path)
        if slide_index >= len(prs.slides):
            return {"status": "error",
                    "error": f"Slide index {slide_index} out of range ({len(prs.slides)} slides)"}
        slide = prs.slides[slide_index]
        notes_slide = slide.notes_slide
        notes_slide.placeholders[1].text = notes
        prs.save(path)
        return {"path": path, "slide_index": slide_index, "notes": notes}

    def apply_theme(self, path: str, theme: str = "dark") -> dict:
        """Apply a simple background theme to all slides.

        Theme options: "dark", "light", "blue". Reuses the palette
        from create_from_slides so decks look consistent.
        """
        from pptx.dml.color import RGBColor as _RGB
        palette = {
            "dark": _RGB(0x1E, 0x1E, 0x2E),
            "light": _RGB(0xFF, 0xFF, 0xFF),
            "blue": _RGB(0x1F, 0x4E, 0x79),
        }
        bg = palette.get(theme)
        if bg is None:
            return {"status": "error",
                    "error": f"Unknown theme: {theme} (options: dark, light, blue)"}
        prs = Presentation(path)
        for slide in prs.slides:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = bg
        prs.save(path)
        return {"path": path, "theme": theme, "slides": len(prs.slides)}

    def get_structure(self, path: str) -> dict:
        """Get full presentation structure."""
        prs = Presentation(path)
        slides = []
        for i, slide in enumerate(prs.slides):
            info = {"index": i, "title": ""}
            if slide.shapes.title:
                info["title"] = slide.shapes.title.text
            info["shape_count"] = len(slide.shapes)
            slides.append(info)
        return {"path": path, "slides": slides, "total": len(slides)}

    def extract_text(self, path: str) -> dict:
        """Extract all text from all slides."""
        prs = Presentation(path)
        content = []
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        if para.text.strip():
                            slide_text.append(para.text)
            content.append({"slide": i, "texts": slide_text})
        return {"path": path, "content": content}

    def create_from_slides(self, path: str, slides: list = None,
                           theme: str = "dark") -> dict:
        """Create a complete presentation from a slides array with rich formatting.

        Each slide can have:
          - type: "title" | "text" | "table" | "kpi" | "chart" (default: "text")
          - title: str
          - subtitle: str (for title slides)
          - content: str (for text slides, lines separated by newline)
          - headers: list[str] (for table slides)
          - rows: list[list[str]] (for table slides)
          - kpis: list[{"label": str, "value": str}] (for kpi slides)
          - chart: {"kind": "bar"|"line"|"pie", "categories": [...], "values": [...]}
          - background: str hex color (optional override)

        Corporate 16:9 layout, Calibri 28pt titles / 18pt body.
        Theme options: "dark" (default), "light", "blue"
        """
        prs = Presentation()
        prs.slide_width = Emu(12192000)  # 16:9 corporate widescreen
        prs.slide_height = Inches(7.5)

        # Theme definitions
        themes = {
            "dark": {
                "bg": RGBColor(0x1E, 0x1E, 0x2E),
                "alt_bg": RGBColor(0x15, 0x15, 0x25),
                "title": RGBColor(0xFF, 0xFF, 0xFF),
                "subtitle": RGBColor(0x00, 0x9E, 0xFF),
                "body": RGBColor(0xE0, 0xE0, 0xE0),
                "muted": RGBColor(0x88, 0x88, 0xAA),
                "accent": RGBColor(0x00, 0x7A, 0xCC),
                "accent2": RGBColor(0x00, 0xCC, 0x88),
                "accent3": RGBColor(0xE8, 0xA0, 0x00),
                "row_even": RGBColor(0x2D, 0x2D, 0x3D),
                "row_odd": RGBColor(0x35, 0x35, 0x45),
                "table_header_bg": RGBColor(0x00, 0x7A, 0xCC),
                "kpi_bg": RGBColor(0x25, 0x25, 0x35),
            },
            "light": {
                "bg": RGBColor(0xFF, 0xFF, 0xFF),
                "alt_bg": RGBColor(0xF5, 0xF5, 0xF5),
                "title": RGBColor(0x1A, 0x1A, 0x2E),
                "subtitle": RGBColor(0x00, 0x7A, 0xCC),
                "body": RGBColor(0x33, 0x33, 0x33),
                "muted": RGBColor(0x66, 0x66, 0x66),
                "accent": RGBColor(0x00, 0x7A, 0xCC),
                "accent2": RGBColor(0x00, 0x96, 0x88),
                "accent3": RGBColor(0xE8, 0xA0, 0x00),
                "row_even": RGBColor(0xF8, 0xF8, 0xF8),
                "row_odd": RGBColor(0xFF, 0xFF, 0xFF),
                "table_header_bg": RGBColor(0x00, 0x7A, 0xCC),
                "kpi_bg": RGBColor(0xF0, 0xF4, 0xF8),
            },
            "blue": {
                "bg": RGBColor(0x1F, 0x4E, 0x79),
                "alt_bg": RGBColor(0x16, 0x3A, 0x5B),
                "title": RGBColor(0xFF, 0xFF, 0xFF),
                "subtitle": RGBColor(0x64, 0xB5, 0xF6),
                "body": RGBColor(0xCF, 0xD8, 0xE6),
                "muted": RGBColor(0x78, 0x90, 0xA5),
                "accent": RGBColor(0x1E, 0x88, 0xE5),
                "accent2": RGBColor(0x4C, 0xAF, 0x50),
                "accent3": RGBColor(0xFF, 0xB3, 0x00),
                "row_even": RGBColor(0x1A, 0x2E, 0x3D),
                "row_odd": RGBColor(0x22, 0x36, 0x45),
                "table_header_bg": RGBColor(0x1E, 0x88, 0xE5),
                "kpi_bg": RGBColor(0x15, 0x25, 0x35),
            },
        }
        t = themes.get(theme, themes["dark"])

        for slide_def in slides or []:
            stype = slide_def.get("type", "text")
            bg_color = t["bg"]
            if "background" in slide_def:
                h = slide_def["background"].lstrip("#")
                bg_color = RGBColor(*tuple(int(h[i:i+2], 16) for i in (0, 2, 4)))

            slide = prs.slides.add_slide(prs.slide_layouts[6])

            # Background
            bg = slide.background
            fill = bg.fill
            fill.solid()
            fill.fore_color.rgb = bg_color

            if stype == "title":
                self._build_title_slide(slide, slide_def, t)
            elif stype == "kpi":
                self._build_kpi_slide(slide, slide_def, t)
            elif stype == "table":
                self._build_table_slide(slide, slide_def, t)
            elif stype == "chart":
                self._build_chart_slide(slide, slide_def, t)
            else:
                self._build_text_slide(slide, slide_def, t)

        prs.save(path)
        return {"path": path, "slides": len(slides or []), "theme": theme}

    def create_deck(self, path: str, title: str = "Corporate Presentation",
                    theme: str = "blue", slides: list = None,
                    notes: list = None) -> dict:
        """Create a corporate 16:9 deck (8 slides by default).

        Default deck: cover, agenda, content, KPI, table, chart,
        content, closing — cover background #1F4E79 with white text.
        Custom slides use the same schema as create_from_slides.
        Optional notes list adds speaker notes per slide index.
        """
        if slides is None:
            slides = [
                {"type": "title", "title": title,
                 "subtitle": "Corporate Report",
                 "background": "#1F4E79"},
                {"type": "text", "title": "Agenda",
                 "content": "1. Overview\n2. Key metrics\n3. Results\n4. Next steps"},
                {"type": "text", "title": "Overview",
                 "content": "Market context and quarterly highlights."},
                {"type": "kpi", "title": "Key Metrics",
                 "kpis": [{"label": "Revenue", "value": "+12%"},
                          {"label": "Customers", "value": "4.8K"},
                          {"label": "Retention", "value": "96%"}]},
                {"type": "table", "title": "Quarterly Results",
                 "headers": ["Metric", "Q3", "Q4"],
                 "rows": [["Revenue", "1.2M", "1.4M"],
                          ["Customers", "4.1K", "4.8K"]]},
                {"type": "chart", "title": "Revenue Trend",
                 "chart": {"kind": "bar",
                           "categories": ["Q1", "Q2", "Q3", "Q4"],
                           "values": [0.9, 1.1, 1.2, 1.4]}},
                {"type": "text", "title": "Next Steps",
                 "content": "Scale what works\nClose open risks"},
                {"type": "title", "title": "Thank You",
                 "subtitle": "Questions & discussion",
                 "background": "#1F4E79"},
            ]
        result = self.create_from_slides(path, slides, theme)
        if notes:
            prs = Presentation(path)
            for i, note in enumerate(notes):
                if note and i < len(prs.slides):
                    prs.slides[i].notes_slide.placeholders[1].text = note
            prs.save(path)
            result["notes"] = sum(1 for n in notes if n)
        return result

    def _add_textbox(self, slide, left, top, width, height, text,
                     size=18, bold=False, color=None, align=PP_ALIGN.LEFT):
        """Helper to add a textbox with formatting."""
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.name = "Calibri"
        p.font.color.rgb = color or RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = align
        return txBox

    def _build_title_slide(self, slide, slide_def, t):
        """Build a title/cover slide."""
        title = slide_def.get("title", "Presentation")
        subtitle = slide_def.get("subtitle", "")
        self._add_textbox(slide, 1, 2.0, 11.3, 1.2, title,
                         size=28, bold=True, color=t["title"], align=PP_ALIGN.CENTER)
        if subtitle:
            self._add_textbox(slide, 1, 3.2, 11.3, 0.6, subtitle,
                             size=18, bold=False, color=t["subtitle"], align=PP_ALIGN.CENTER)
        content = slide_def.get("content", "")
        if content:
            self._add_textbox(slide, 1, 4.0, 11.3, 1.0, content,
                             size=18, bold=False, color=t["muted"], align=PP_ALIGN.CENTER)

    def _build_text_slide(self, slide, slide_def, t):
        """Build a text/content slide."""
        title = slide_def.get("title", "")
        if title:
            self._add_textbox(slide, 0.5, 0.3, 12.3, 0.8, title,
                             size=28, bold=True, color=t["title"])
        content = slide_def.get("content", "")
        if content:
            lines = content.split("\n")
            y = 1.4
            for line in lines:
                if line.strip():
                    display = line.strip()
                    is_bullet = display.startswith("•") or display.startswith("-")
                    color = t["body"] if not is_bullet else t["subtitle"]
                    self._add_textbox(slide, 0.8 if is_bullet else 0.5,
                                     y, 11.8, 0.45, display,
                                     size=18, bold=not is_bullet, color=color)
                    y += 0.4

    def _build_table_slide(self, slide, slide_def, t):
        """Build a table slide."""
        title = slide_def.get("title", "")
        if title:
            self._add_textbox(slide, 0.5, 0.3, 12.3, 0.6, title,
                             size=28, bold=True, color=t["title"])

        headers = slide_def.get("headers", [])
        rows = slide_def.get("rows", [])
        if not headers and not rows:
            return

        nrows = (1 if headers else 0) + len(rows)
        ncols = len(headers) if headers else (len(rows[0]) if rows else 1)
        col_width = 12.3 / max(ncols, 1)
        row_height = 0.5
        top_offset = 1.3 if title else 0.5

        table_shape = slide.shapes.add_table(
            nrows, ncols,
            Inches(0.5), Inches(top_offset),
            Inches(12.3), Inches(row_height * max(nrows, 1))
        )
        table = table_shape.table

        # Header row
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = str(h)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(18)
                p.font.bold = True
                p.font.name = "Calibri"
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                p.alignment = PP_ALIGN.CENTER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = t["table_header_bg"]

        # Data rows
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = table.cell(ri + (1 if headers else 0), ci)
                cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(18)
                    p.font.name = "Calibri"
                    p.font.color.rgb = t["body"]
                    p.alignment = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = t["row_even"] if ri % 2 == 0 else t["row_odd"]

    def _build_kpi_slide(self, slide, slide_def, t):
        """Build a KPI/metrics dashboard slide."""
        title = slide_def.get("title", "")
        if title:
            self._add_textbox(slide, 0.5, 0.3, 12.3, 0.6, title,
                             size=28, bold=True, color=t["title"])

        kpis = slide_def.get("kpis", [])
        if not kpis:
            return

        cols = min(4, len(kpis))
        rows_count = (len(kpis) + cols - 1) // cols
        box_w = 11.0 / cols
        box_h = 1.3

        start_y = 1.5
        start_x = 0.8

        accents = [t["accent"], t["accent2"], t["accent3"], t["subtitle"]]

        for i, kpi in enumerate(kpis):
            col = i % cols
            row = i // cols
            x = start_x + col * box_w
            y = start_y + row * (box_h + 0.3)
            accent = accents[col % len(accents)]

            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(y), Inches(box_w - 0.3), Inches(box_h)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = t["kpi_bg"]
            shape.line.color.rgb = accent
            shape.line.width = Pt(2)

            tf = shape.text_frame
            tf.word_wrap = True
            p1 = tf.paragraphs[0]
            p1.text = str(kpi.get("value", ""))
            p1.font.size = Pt(28)
            p1.font.bold = True
            p1.font.name = "Calibri"
            p1.font.color.rgb = accent
            p1.alignment = PP_ALIGN.CENTER

            p2 = tf.add_paragraph()
            p2.text = str(kpi.get("label", ""))
            p2.font.size = Pt(18)
            p2.font.name = "Calibri"
            p2.font.color.rgb = t["muted"]
            p2.alignment = PP_ALIGN.CENTER

    def _build_chart_slide(self, slide, slide_def, t):
        """Build a chart slide (matplotlib PNG embedded)."""
        from python_engine.charts import ChartEngine
        import tempfile

        title = slide_def.get("title", "")
        if title:
            self._add_textbox(slide, 0.5, 0.3, 12.3, 0.6, title,
                              size=28, bold=True, color=t["title"])

        spec = slide_def.get("chart", {}) or {}
        categories = spec.get("categories", [])
        values = spec.get("values", [])
        if not categories or not values:
            return

        kind = spec.get("kind", "bar")
        engine = ChartEngine()
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            img_path = tmp.name
        try:
            if kind == "pie":
                engine.pie_chart(labels=categories, values=values,
                                 title="", path=img_path)
            elif kind == "line":
                engine.line_chart(x_data=categories, y_series=values,
                                  title="", path=img_path)
            else:
                engine.bar_chart(categories=categories, values=values,
                                 title="", color="#1F4E79", path=img_path)
            slide.shapes.add_picture(img_path, Inches(2.15), Inches(1.4),
                                     width=Inches(9.0))
        finally:
            os.unlink(img_path)
